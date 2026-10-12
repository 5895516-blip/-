# -*- coding: utf-8 -*-
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from style import *
from style import _fix_grid

OUT = "КП_AI_в_кадре_видеобаннер_для_клиники.docx"
doc = new_doc()


def three_price_boxes(doc, items, gap=Cm(0.4)):
    w = (CONTENT_W - 2 * gap) / 3
    widths = [w, gap, w, gap, w]
    fills = ["3B3C36", TEAL_HEX, "101010"]
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    k = 0
    for i, c in enumerate(t.rows[0].cells):
        c.width = Emu(int(widths[i])); no_borders(c)
        if i % 2:
            cell_margins(c, 0, 0, 0, 0); continue
        data = items[k]; shade(c, fills[k]); k += 1
        cell_margins(c, 160, 160, 120, 120)
        for j, (txt, sz, bold, italic) in enumerate(((data[0].upper(), 9, False, False),
                                                     (data[1], 13, True, False),
                                                     (data[2], 9, False, True))):
            p = c.paragraphs[0] if j == 0 else c.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15
            run(p, txt, size=sz, bold=bold, italic=italic, color=WHITE if bold else PALE,
                spacing=0.6 if j == 0 else None)
    _fix_grid(t, widths); return t


# --- обложка
p = para(doc, align=AL.CENTER, space_after=8)
p.add_run().add_picture("logo.png", width=Cm(5.8))
para(doc, "AI В КАДРЕ", size=10, bold=True, color=TEAL, align=AL.CENTER, space_after=2, spacing=1.2)
para(doc, "студия ИИ-видеопродакшна  ·  Москва  ·  2026", size=10, italic=True,
     color=GREY, align=AL.CENTER, space_after=20)
para(doc, "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", size=22, bold=True, align=AL.CENTER, space_after=14)
para(doc, "Видеобаннер для клиники · 0:10", size=14, bold=True, color=TEAL, align=AL.CENTER, space_after=6)
para(doc, "для Артура  ·  два варианта", size=13, italic=True, align=AL.CENTER, space_after=4)
para(doc, "7 рабочих дней  ·  предоплата 50%  ·  скидка 10% до 15 октября", size=10, italic=True,
     color=GREY, align=AL.CENTER, space_after=10)
rule(doc, space_before=6, space_after=14)
two_price_boxes(doc,
                ("Стандарт · 2 правки", "15 300 ₽", "вместо 17 000 ₽ · скидка 10% до 15 октября"),
                ("Расширенный · 4 правки", "18 000 ₽", "вместо 20 000 ₽ · скидка 10% до 15 октября"))
para(doc, "", space_after=14)
stat_row(doc, [("600+", "проектов в ИИ-видео"), ("200 млн+", "просмотров у наших роликов"),
               ("200 000+", "подписчиков в соцсетях")])
para(doc, "", space_after=10)
para(doc, "Цены без НДС — студия работает на УСН.  ·  +7 (968) 589-55-16  ·  aivkadre@yandex.ru",
     size=9.5, italic=True, color=GREY, align=AL.CENTER)
page_break(doc)

# --- задача
h1(doc, "Что входит в видеобаннер", space_before=0)
para(doc, "Артур, добрый день! Как и обсуждали, направляем предложение по видеобаннеру для вашей клиники. Короткий, "
          "аккуратный ролик на 10 секунд, который вызывает доверие и запоминается. Предлагаем два варианта — они "
          "отличаются количеством правок.")
para(doc, "За 10 секунд ролик должен показать главное: уровень клиники, заботу о пациенте и то, куда обращаться. "
          "Делаем его чистым и спокойным по картинке — так, как ожидают от медицины.")
kv_table(doc, [
    ("Сценарий", "Раскладываем 10 секунд по кадрам под вашу услугу или акцию"),
    ("Визуал", "Светлые современные кабинеты, врач и пациент, оборудование — по вашему брендбуку"),
    ("Графика", "Логотип, ключевая фраза или оффер, контакты клиники"),
    ("Звук", "Лицензионная фоновая музыка"),
    ("Предупреждение", "Обязательная для рекламы медицинских услуг плашка о противопоказаниях"),
    ("Срок", "7 рабочих дней"),
    ("Правки", "2 — в «Стандарте», 4 — в «Расширенном»"),
])
para(doc, "Правка — одно конкретное изменение: заменить план, поправить надпись, сдвинуть склейку.",
     size=9.5, italic=True, color=GREY, space_before=4)

h1(doc, "Как работаем")
for title, body in [
    ("Бриф.", "Уточняем услугу или акцию, тон и где будет показан баннер."),
    ("Сценарий и стилевые кадры.", "Утверждаете, как будет выглядеть ролик."),
    ("Производство.", "Генерация, анимация, монтаж, графика, музыка."),
    ("Правки и сдача.", "Вносим правки и отдаём готовый файл."),
]:
    p = para(doc, space_after=5)
    run(p, title + " ", bold=True); run(p, body, color=BODY)

h1(doc, "Стоимость и оплата")
keep_table_together(grid_table(doc, ["Вариант", "Цена", "Со скидкой 10%", "Предоплата 50%"], [
    ("Стандарт · 2 правки", "17 000 ₽", "15 300 ₽", "7 650 ₽"),
    ("Расширенный · 4 правки", "20 000 ₽", "18 000 ₽", "9 000 ₽"),
], [Cm(5.7), Cm(3.4), Cm(3.7), Cm(3.7)]))
para(doc, "Скидка 10% действует при оплате до 15 октября включительно. Предоплата 50% — при старте, остаток — после "
          "утверждения готового ролика. Цена без НДС: студия на УСН, сверх согласованного объёма счетов не выставляем.",
     size=9.5, italic=True, color=GREY, space_before=6)
sub(doc, "Если планируете больше роликов", space_before=10)
para(doc, "При заказе 3–5 видео сразу с предоплатой готовы рассмотреть дополнительную скидку 10%. Если нужно больше "
          "роликов — напишите, сколько и каких: просчитаем индивидуально.", color=BODY)

h1(doc, "Реквизиты")
keep_table_together(kv_table(doc, [
    ("Исполнитель", "Студия AI в кадре · ИП Тальян Антон Дмитриевич"),
    ("ИНН / ОГРНИП", "774349397920 · 325774600866130"),
    ("Расчётный счёт", "40802810700009144244, АО «ТИНЬКОФФ БАНК», г. Москва"),
    ("Корр. счёт / БИК", "30101810145250000974 · 044525974"),
    ("Контакты", "+7 (968) 589-55-16 (телефон / Telegram) · aivkadre@yandex.ru · aivkadre.ru"),
]))

p = para(doc, space_before=6, space_after=0)
run(p, "Готовы стартовать. ", bold=True, color=TEAL)
run(p, "Артур, напишите, какой вариант вам ближе, — пришлём счёт на предоплату. Успейте до 15 октября, чтобы "
       "зафиксировать скидку 10%.", color=BODY)

doc.save(OUT)
print("saved", OUT)
