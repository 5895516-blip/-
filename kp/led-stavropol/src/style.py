# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x0E, 0x8F, 0x7D); TEAL_HEX = "0E8F7D"
TEAL_BRIGHT = "16B8A5"; TEAL_LIGHT = "E6F7F5"; BORDER = "E2E8E6"
INK = RGBColor(0x10, 0x10, 0x10); BODY = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x6B, 0x6B, 0x6B); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xE6, 0xF7, 0xF5)
FONT = "Arial"; CONTENT_W = Cm(16.5)


def _set_font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def run(p, text, size=11, bold=False, italic=False, color=INK, caps=False, spacing=None):
    r = p.add_run(text.upper() if caps else text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = color
    _set_font(r)
    if spacing:
        el = OxmlElement("w:spacing"); el.set(qn("w:val"), str(int(spacing * 20)))
        r._element.get_or_add_rPr().append(el)
    return r


def para(doc, text="", size=11, bold=False, italic=False, color=INK, align=None,
         space_before=0, space_after=6, line=1.25, caps=False, spacing=None, container=None):
    host = container if container is not None else doc
    p = host.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size, bold, italic, color, caps, spacing)
    return p


def bottom_border(p, color=TEAL_HEX, size=12):
    bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    bd.append(b); p._p.get_or_add_pPr().append(bd)


def keep_next(p):
    p._p.get_or_add_pPr().append(OxmlElement("w:keepNext")); return p


def h1(doc, text, space_before=20):
    p = para(doc, text, size=18, bold=True, color=TEAL, caps=True,
             space_before=space_before, space_after=10, line=1.1)
    bottom_border(p, TEAL_HEX, 12); keep_next(p); return p


def sub(doc, text, space_before=8, space_after=8):
    p = para(doc, text, size=12, bold=True, color=TEAL,
             space_before=space_before, space_after=space_after)
    keep_next(p); return p


def rule(doc, color=TEAL_BRIGHT, space_before=10, space_after=10, size=18):
    p = para(doc, "", space_before=space_before, space_after=space_after)
    bottom_border(p, color, size); return p


def shade(cell, hex_color):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(sh)


def cell_borders(cell, color=BORDER, size=4):
    tcB = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + e)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        tcB.append(el)
    cell._tc.get_or_add_tcPr().append(tcB)


def no_borders(cell):
    tcB = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + e); el.set(qn("w:val"), "nil"); tcB.append(el)
    cell._tc.get_or_add_tcPr().append(tcB)


def cell_margins(cell, top=100, bottom=100, left=140, right=140):
    m = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + name); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
        m.append(el)
    cell._tc.get_or_add_tcPr().append(m)


def vcenter(cell):
    va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center")
    cell._tc.get_or_add_tcPr().append(va)


def _prep_cell(cell, width):
    cell.width = Emu(int(width))
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2


def _fix_grid(t, widths):
    tbl = t._tbl; tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(Emu(int(w)).twips)))
        grid.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, w in zip(row.cells, widths):
            c.width = Emu(int(w))


def keep_table_together(t):
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                keep_next(p)
    return t


def kv_table(doc, rows, w1=Cm(4.6), stripe=True, label_size=10.5, body_size=10.5):
    w2 = CONTENT_W - w1
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    for i, (label, body) in enumerate(rows):
        cells = t.add_row().cells
        bg = TEAL_LIGHT if (stripe and i % 2 == 0) else "FFFFFF"
        for c, w in zip(cells, (w1, w2)):
            _prep_cell(c, w); shade(c, bg); cell_borders(c); cell_margins(c)
        run(cells[0].paragraphs[0], label, size=label_size, bold=True, color=TEAL)
        run(cells[1].paragraphs[0], body, size=body_size, color=BODY)
    _fix_grid(t, [w1, w2]); return t


def grid_table(doc, header, rows, widths, stripe=True, size=10.5):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    if header:
        cells = t.add_row().cells
        for c, w, txt in zip(cells, widths, header):
            _prep_cell(c, w); shade(c, TEAL_HEX); cell_borders(c, TEAL_HEX)
            cell_margins(c); vcenter(c)
            run(c.paragraphs[0], txt, size=size - 0.5, bold=True, color=WHITE)
        cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        bg = TEAL_LIGHT if (stripe and i % 2 == 0) else "FFFFFF"
        for j, (c, w, txt) in enumerate(zip(cells, widths, r)):
            _prep_cell(c, w); shade(c, bg); cell_borders(c); cell_margins(c)
            run(c.paragraphs[0], txt, size=size, bold=(j == 0),
                color=TEAL if j == 0 else BODY)
    _fix_grid(t, widths); return t


def price_box(doc, label, amount, sub_, sub2=None, width=CONTENT_W, fill=TEAL_HEX):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    c = t.rows[0].cells[0]; c.width = Emu(int(width))
    shade(c, fill); no_borders(c); cell_margins(c, 200, 200, 220, 220)
    for txt, sz, bold, italic, col in ((label.upper(), 10, False, False, PALE),
                                       (amount, 15, True, False, WHITE),
                                       (sub_, 10, False, True, PALE)) + \
                                      (((sub2, 10, False, True, PALE),) if sub2 else ()):
        p = c.paragraphs[0] if txt == label.upper() else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
        run(p, txt, size=sz, bold=bold, italic=italic, color=col,
            spacing=0.8 if col is PALE and not italic else None)
    _fix_grid(t, [width]); return t


def two_price_boxes(doc, left, right, gap=Cm(0.5)):
    w = (CONTENT_W - gap) / 2
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    widths = [w, gap, w]
    for i, c in enumerate(t.rows[0].cells):
        c.width = Emu(int(widths[i])); no_borders(c)
        if i == 1:
            cell_margins(c, 0, 0, 0, 0); continue
        data = left if i == 0 else right
        shade(c, TEAL_HEX); cell_margins(c, 180, 180, 160, 160)
        for k, (txt, sz, bold, italic) in enumerate(((data[0].upper(), 9.5, False, False),
                                                     (data[1], 14, True, False),
                                                     (data[2], 9.5, False, True))):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15
            run(p, txt, size=sz, bold=bold, italic=italic,
                color=WHITE if bold else PALE, spacing=0.8 if k == 0 else None)
    _fix_grid(t, widths); return t


def stat_row(doc, items):
    n = len(items); w = CONTENT_W / n
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, (num, cap) in zip(t.rows[0].cells, items):
        c.width = Emu(int(w)); no_borders(c); cell_margins(c, 60, 60, 80, 80)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run(p, num, size=17, bold=True, color=TEAL)
        p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        run(p2, cap, size=9.5, color=GREY)
    _fix_grid(t, [w] * n); return t


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_height = Cm(29.7); s.page_width = Cm(21.0)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.3)
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(11); st.font.color.rgb = INK
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25
    return doc
