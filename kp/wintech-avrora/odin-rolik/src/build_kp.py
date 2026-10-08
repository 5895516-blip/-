# -*- coding: utf-8 -*-
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from style import *
from style import _fix_grid

OUT = "КП_AI_в_кадре_WINTECH_Полёт_триумфатора.docx"
doc = new_doc()


def three_price_boxes(doc, items, gap=Cm(0.4)):
    w = (CONTENT_W - 2 * gap) / 3
    widths = [w, gap, w, gap, w]
    fills = ["3B3C36", TEAL_HEX, "101010"]
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    k = 0
    for i, c in enumerate(t.rows[0].cells):
        c.width = Emu(int(widths[i])); no_borders(c)
        if i % 2:
            cell_margins(c, 0, 0, 0, 0); continue
        data = items[k]; shade(c, fills[k]); k += 1
        cell_margins(c, 160, 160, 120, 120)
        for j, (txt, sz, bold, italic) in enumerate(((data[0].upper(), 9, False, False),
                                                     (data[1], 13, True, False),
                                                     (data[2], 9, False, True))):
            p = c.paragraphs[0] if j == 0 else c.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15
            run(p, txt, size=sz, bold=bold, italic=italic, color=WHITE if bold else PALE,
                spacing=0.6 if j == 0 else None)
    _fix_grid(t, widths); return t


# --- обложка
p = para(doc, align=AL.CENTER, space_after=8)
p.add_run().add_picture("logo.png", width=Cm(5.8))
para(doc, "AI В КАДРЕ", size=10, bold=True, color=TEAL, align=AL.CENTER, space_after=2, spacing=1.2)
para(doc, "студия ИИ-видеопродакшна  ·  Москва  ·  2026", size=10, italic=True,
     color=GREY, align=AL.CENTER, space_after=20)
para(doc, "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", size=22, bold=True, align=AL.CENTER, space_after=14)
para(doc, "Полёт триумфатора · WINTECH · 0:15", size=14, bold=True, color=TEAL, align=AL.CENTER, space_after=6)
para(doc, "для AVRORA  ·  рекламный ролик", size=13, italic=True, align=AL.CENTER, space_after=4)
para(doc, "16:9  ·  15 секунд  ·  5 сцен  ·  7 рабочих дней", size=10, italic=True,
     color=GREY, align=AL.CENTER, space_after=10)
rule(doc, space_before=6, space_after=14)
two_price_boxes(doc,
                ("Ролик · 15 секунд", "35 000 ₽", "7 рабочих дней · 3 правки"),
                ("Со скидкой 7%", "32 550 ₽", "скидка на первый заказ"))
para(doc, "", space_after=14)
stat_row(doc, [("600+", "проектов в ИИ-видео"), ("200 млн+", "просмотров у наших роликов"),
               ("200 000+", "подписчиков в соцсетях")])
para(doc, "", space_after=10)
para(doc, "Цены без НДС — студия работает на УСН. Оплата 50 / 50.", size=9.5, italic=True,
     color=GREY, align=AL.CENTER, space_after=4)
para(doc, "+7 (968) 589-55-16  ·  aivkadre@yandex.ru  ·  aivkadre.ru", size=10, color=GREY, align=AL.CENTER)
page_break(doc)

# --- задача
h1(doc, "Что мы поняли из задачи", space_before=0)
para(doc, "Кирилл, добрый день! Да, стартуем с одного ролика: «Полёт триумфатора», 15 секунд, 5 сцен, формат 16:9. "
          "С учётом скидки 7% на первый заказ стоимость — 32 550 ₽.")
para(doc, "По ночному небу несётся тройка белоснежных коней. Она пролетает над городами и посёлками, над особняками "
          "и небоскрёбами, над югом и севером. Там, где пролетает тройка, окна WINTECH вспыхивают, и люди подходят "
          "посмотреть. Финал — слоган «Мир доверяет окнам WINTECH».")
sub(doc, "Примерная раскладка 5 сцен")
kv_table(doc, [
    ("Сцена 1", "Ночное небо над городом, появляется тройка со светящимся шлейфом."),
    ("Сцена 2", "Пролёт над частными особняками, окна вспыхивают."),
    ("Сцена 3", "Пролёт над небоскрёбами мегаполиса."),
    ("Сцена 4", "Юг и север, люди подходят к окнам и смотрят вслед."),
    ("Сцена 5", "Финал: слоган «Мир доверяет окнам WINTECH»."),
])
para(doc, "Финальную раскладку определит ваш сценарий — подстроимся под него.", size=9.5, italic=True,
     color=GREY, space_before=6)

h1(doc, "Как устроена работа")
for title, body in [
    ("Стилевые кадры.", "Тройка, ночной город, свечение окон. Можно начать по эскизам, пока вы дописываете сценарий."),
    ("Раскадровка.", "Делим сценарий на 5 сцен по секундам."),
    ("Генерация и анимация.", "Сцены, движение камеры, переходы. Каждый план в нескольких дублях, берём лучший."),
    ("Музыка и монтаж.", "Мелодия со слоганом, финальная плашка, цвет."),
    ("Правки и сдача.", "Черновой монтаж, правки, финальный файл."),
]:
    p = para(doc, space_after=5)
    run(p, title + " ", bold=True); run(p, body, color=BODY)

h1(doc, "Сроки")
keep_table_together(grid_table(doc, ["Этап", "Результат", "Дней"], [
    ("Стилевые кадры и раскадровка", "Кадры тройки и города, план сцен", "2"),
    ("Генерация и анимация", "Все 5 сцен", "3"),
    ("Музыка и монтаж", "Черновой монтаж", "1"),
    ("Правки и финал", "Готовый ролик", "1"),
    ("Итого", "", "7"),
], [Cm(6.0), Cm(8.0), Cm(2.5)], size=10))
para(doc, "Время на ваши согласования в срок не входит: обычно это по дню на стилевые кадры и на черновик.",
     size=9.5, italic=True, color=GREY, space_before=6)

# --- стоимость
h1(doc, "Что входит")
kv_table(doc, [
    ("Хронометраж", "15 с, 5 сцен"),
    ("Формат", "16:9, Full HD"),
    ("Сценарий", "Ваш, раскладываем по сценам и секундам"),
    ("Стилевые кадры", "2 кадра до старта генерации"),
    ("Визуал", "Тройка белых коней, ночные города, свечение окон WINTECH"),
    ("Звук", "Мелодия со слоганом"),
    ("Графика", "Слоган, финальная плашка"),
    ("Правки", "3 правки"),
])

h1(doc, "Сколько выходит")
keep_table_together(grid_table(doc, ["", "Сумма"], [
    ("Ролик «Полёт триумфатора», 15 с", "35 000 ₽"),
    ("Скидка 7% на первый заказ", "−2 450 ₽"),
    ("Аванс при старте, 50%", "16 275 ₽"),
    ("После утверждения финала, 50%", "16 275 ₽"),
    ("Итого", "32 550 ₽"),
], [Cm(11.5), Cm(5.0)]))
para(doc, "Правка — одно конкретное изменение: заменить план, поправить надпись, сдвинуть склейку. "
          "Цена без НДС — студия работает на УСН. Дополнительных счетов сверх согласованного объёма не выставляем.",
     size=9.5, italic=True, color=GREY, space_before=6)

h1(doc, "Что нужно от вас для старта")
kv_table(doc, [
    ("Обязательно", "Сценарий ролика: 5 сцен, слова мелодии, слоган. Фирменные цвета и шрифты WINTECH. "
                    "Фото профилей и окон WINTECH. Контакт, который утверждает этапы."),
    ("Полезно", "Эскизы из презентации в исходном качестве. Музыкальные референсы. Где будет размещён ролик."),
])

h1(doc, "Реквизиты")
keep_table_together(kv_table(doc, [
    ("Исполнитель", "Студия AI в кадре · ИП Тальян Антон Дмитриевич"),
    ("ИНН / ОГРНИП", "774349397920 · 325774600866130"),
    ("Расчётный счёт", "40802810700009144244, АО «ТИНЬКОФФ БАНК», г. Москва"),
    ("Корр. счёт / БИК", "30101810145250000974 · 044525974"),
    ("Контакты", "+7 (968) 589-55-16 (телефон / Telegram) · aivkadre@yandex.ru · aivkadre.ru"),
]))

p = para(doc, space_before=4, space_after=0)
run(p, "Готовы стартовать. ", bold=True, color=TEAL)
run(p, "Кирилл, пришлём договор и счёт на аванс 16 275 ₽. Стилевые кадры тройки покажем через 2 рабочих дня "
       "после старта.", color=BODY)

doc.save(OUT)
print("saved", OUT)
